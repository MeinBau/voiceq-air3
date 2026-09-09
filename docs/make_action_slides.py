"""조치(Action) 계층 발표 슬라이드 생성기 — 팀 덱 본문 스타일(흰 배경·네이비).

finetune/make_deck_pptx.py 와는 팔레트가 다르다. 저쪽은 어두운 지휘소 톤이고,
이 파일은 팀이 실제로 쓰고 있는 발표 덱 본문 슬라이드(흰 배경 + 네이비 + 상단
breadcrumb + 하단 푸터) 형식을 따른다. 그 덱에 그대로 끼워 넣기 위한 것이다.

내용은 modules/actions.py 로 구현한 "판단과 실행의 분리"다. 수치·문구는 전부 실제
코드와 tests/test_actions.py 검증 결과에서 가져왔다.

사용:
    python docs/make_action_slides.py                    # docs/voicecue_조치계층.pptx
    python docs/make_action_slides.py --out 경로.pptx
    python docs/make_action_slides.py --start-page 24    # 푸터 쪽번호 시작값
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# ---------------------------------------------------------------------
# 팔레트 — 팀 덱 본문 슬라이드에서 뽑은 값
# ---------------------------------------------------------------------
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x14, 0x33, 0x6B)          # 제목·배너·강조 박스
NAVY_DEEP = RGBColor(0x0D, 0x22, 0x4A)     # 번호 뱃지
BLUE = RGBColor(0x2E, 0x6D, 0xB4)          # 보조 강조
BLUE_PALE = RGBColor(0xE9, 0xF0, 0xFA)     # 연한 채움
PILL_BG = RGBColor(0xEE, 0xF2, 0xF7)       # 비활성 단계 pill
INK = RGBColor(0x1A, 0x1A, 0x1A)
DIM = RGBColor(0x6B, 0x76, 0x84)
FAINT = RGBColor(0x9A, 0xA5, 0xB1)
LINE = RGBColor(0xD8, 0xE0, 0xE9)
OK = RGBColor(0x1E, 0x8A, 0x5F)
WARN = RGBColor(0xC2, 0x7A, 0x1B)

# 목업(어두운 화면) 톤
MOCK_BG = RGBColor(0x18, 0x22, 0x30)
MOCK_PANEL = RGBColor(0x22, 0x2E, 0x3E)
MOCK_INK = RGBColor(0xE4, 0xEE, 0xF7)
MOCK_DIM = RGBColor(0x8F, 0xA8, 0xC2)

SLIDE_W = Emu(12192000)   # 33.87cm
SLIDE_H = Emu(6858000)    # 19.05cm

FONT = "맑은 고딕"
MONO = "Consolas"

FOOTER_LEFT = "대한민국 공군 · VOICE-CUE · 팀 작비스"
TOTAL_PAGES = 37


def cm(v: float) -> Emu:
    return Emu(int(v * 360000))


class Deck:
    """팀 덱 본문 슬라이드 한 장을 만드는 데 필요한 것만."""

    def __init__(self, start_page: int) -> None:
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.page = start_page

    # ---------- 원시 도형 ----------
    def _text(self, slide, x, y, w, h, runs, *, size=12, color=INK, bold=False,
              font=FONT, align=PP_ALIGN.LEFT, spacing=1.3,
              anchor=MSO_ANCHOR.TOP, space_after=0):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        items = runs if isinstance(runs, list) else [runs]
        first = True
        for item in items:
            text, opts = item if isinstance(item, tuple) else (item, {})
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = opts.get("align", align)
            p.line_spacing = opts.get("spacing", spacing)
            if opts.get("space_before"):
                p.space_before = Pt(opts["space_before"])
            if space_after:
                p.space_after = Pt(space_after)
            run = p.add_run()
            run.text = text
            f = run.font
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.name = opts.get("font", font)
            f.color.rgb = opts.get("color", color)
        return box

    def _rect(self, slide, x, y, w, h, fill, *, line=None, radius=None):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = slide.shapes.add_shape(shape_type, x, y, w, h)
        if radius is not None:
            # adjustment 0 = 직각, 클수록 둥글다. 높이 대비 비율로 들어간다.
            sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(0.75)
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        return sh

    # ---------- 슬라이드 골격 ----------
    def slide(self, breadcrumb: str, title: str, banner: str):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = WHITE

        self._text(s, cm(1.5), cm(0.85), cm(30), cm(0.6), breadcrumb,
                   size=10.5, color=NAVY, bold=True)
        self._text(s, cm(1.5), cm(1.6), cm(30), cm(1.4), title,
                   size=27, color=NAVY, bold=True)

        bar = self._rect(s, cm(1.5), cm(3.35), cm(30.9), cm(1.05), NAVY)
        tf = bar.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = cm(0.5)
        tf.margin_right = cm(0.5)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = banner
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = FONT
        run.font.color.rgb = WHITE

        # 푸터
        self._text(s, cm(1.5), cm(17.75), cm(20), cm(0.5), FOOTER_LEFT,
                   size=9, color=FAINT)
        self._text(s, cm(27), cm(17.75), cm(5.4), cm(0.5),
                   f"{self.page} / {TOTAL_PAGES}", size=9, color=FAINT,
                   align=PP_ALIGN.RIGHT)
        self.page += 1
        return s

    # ---------- 구성 요소 ----------
    def pills(self, slide, y, steps, active_idx):
        """상단 단계 표시줄. steps: [(번호, 라벨)]"""
        n = len(steps)
        gap = cm(0.3)
        total = cm(30.9)
        w = Emu(int((total - gap * (n - 1)) / n))
        x = cm(1.5)
        for i, (num, label) in enumerate(steps):
            on = i == active_idx
            sh = self._rect(slide, x, y, w, cm(0.95),
                            NAVY if on else PILL_BG,
                            line=None if on else LINE, radius=0.35)
            tf = sh.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = cm(0.2)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = f"{num}. {label}"
            r.font.size = Pt(10.5)
            r.font.bold = on
            r.font.name = FONT
            r.font.color.rgb = WHITE if on else DIM
            x = Emu(x + w + gap)

    def section_label(self, slide, x, y, text, w=15.1):
        # 폭을 인자로 받는다 — 우측 열 라벨은 좌측보다 길어질 수 있는데(파일명·항목수),
        # 고정 폭이면 두 줄로 접히면서 아래 본문과 겹친다.
        self._text(slide, x, y, cm(w), cm(0.6), text, size=11, color=BLUE, bold=True)

    def steps(self, slide, x, y, w, items):
        """좌측 번호 박스 목록. items: [(제목, 부제)]"""
        h = cm(1.75)
        gap = cm(0.28)
        cur = y
        for i, (title, sub) in enumerate(items, start=1):
            box = self._rect(slide, x, cur, w, h, NAVY, radius=0.12)
            # 번호 뱃지
            badge = self._rect(slide, Emu(x + cm(0.35)), Emu(cur + cm(0.45)),
                               cm(0.85), cm(0.85), NAVY_DEEP, radius=0.5)
            btf = badge.text_frame
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            br = bp.add_run()
            br.text = str(i)
            br.font.size = Pt(12)
            br.font.bold = True
            br.font.name = FONT
            br.font.color.rgb = WHITE

            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = cm(1.45)
            tf.margin_right = cm(0.3)
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = title
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = WHITE
            p2 = tf.add_paragraph()
            p2.line_spacing = 1.15
            r2 = p2.add_run()
            r2.text = sub
            r2.font.size = Pt(9)
            r2.font.name = FONT
            r2.font.color.rgb = RGBColor(0xB8, 0xC9, 0xE0)
            cur = Emu(cur + h + gap)
        return cur

    def bullets(self, slide, x, y, w, items):
        runs = []
        for i, t in enumerate(items):
            runs.append((f"·  {t}", {"space_before": 0 if i == 0 else 7}))
        self._text(slide, x, y, w, cm(4), runs, size=11.5, color=INK, spacing=1.3)


def build(deck: Deck) -> None:
    # =================================================================
    # 슬라이드 1 — 왜 분리했는가
    # =================================================================
    s = deck.slide(
        "04 핵심 기술 및 구현  ·  B. 개발 과정 및 구현 결과  ·  조치 계층",
        "판단과 실행의 분리 — 물리 장비가 붙을 자리",
        "기존: 판단 결과를 화면 상태에 직접 대입 → 실패를 표현할 수 없고, PTZ·Video Wall을 끼워 넣을 자리가 없음",
    )
    deck.pills(s, cm(4.75), [
        ("1", "발언 입력"), ("2", "상황 판단"), ("3", "조치 생성"),
        ("4", "장비 실행"), ("5", "결과 보고"),
    ], active_idx=2)

    deck.section_label(s, cm(1.5), cm(6.2), "문제")
    deck.steps(s, cm(1.5), cm(6.85), cm(14.2), [
        ("대입은 실패할 수 없다",
         "cop_layout = layout — “띄우라 했는데 안 떴다”를 표현할 방법이 없음"),
        ("장비 연동 지점이 없다",
         "판단 코드 한복판에 하드웨어 호출을 끼워 넣는 것 외 방법 없음"),
        ("조치가 하나로 뭉개진다",
         "화면 표출과 카메라 지향은 서로 다른 장비의 조치인데 구분 불가"),
    ])

    deck.section_label(s, cm(17.3), cm(6.2), "구조")
    deck.bullets(s, cm(17.3), cm(6.9), cm(15.1), [
        "판단 결과 → Action 목록으로 변환 (무엇을·왜)",
        "ActionBus가 조치를 지원하는 Actuator에게 위임",
        "결과는 수행 / 미연동 / 실패 세 가지로 구분해 보고",
        "actions.py는 streamlit 비의존 — 실행 주체가 화면일 필요 없음",
    ])

    code = deck._rect(s, cm(17.3), cm(10.4), cm(15.1), cm(6.3), MOCK_BG, radius=0.05)
    ctf = code.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_top = cm(0.45)
    ctf.margin_right = cm(0.3)
    lines = [
        ("@dataclass(frozen=True)", MOCK_DIM),
        ("class Action:", MOCK_INK),
        ("    kind: str        # 화면구성 | 카메라지향", MOCK_DIM),
        ("    payload: dict", MOCK_INK),
        ("    reason: str      # 근거 없는 조치는 만들지 않는다", MOCK_DIM),
        ("", MOCK_INK),
        ("class Actuator(Protocol):", MOCK_INK),
        ("    def supports(self, kind) -> bool: ...", MOCK_INK),
        ("    def execute(self, action) -> ActionResult: ...", MOCK_INK),
    ]
    for i, (text, color) in enumerate(lines):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.line_spacing = 1.35
        r = p.add_run()
        r.text = text
        r.font.size = Pt(10)
        r.font.name = MONO
        r.font.color.rgb = color

    # =================================================================
    # 슬라이드 2 — 하나의 판단이 여러 장비 조치로
    # =================================================================
    s = deck.slide(
        "04 핵심 기술 및 구현  ·  B. 개발 과정 및 구현 결과  ·  조치 계층",
        "하나의 판단 → 서로 다른 장비의 조치로 분기",
        "“북서 CCTV를 띄운다”는 결정은 실제 지휘소에서 ① 벽면 표출 ② 그 카메라 지향, 두 장비에 대한 두 조치다",
    )
    deck.pills(s, cm(4.75), [
        ("1", "발언 입력"), ("2", "상황 판단"), ("3", "조치 생성"),
        ("4", "장비 실행"), ("5", "결과 보고"),
    ], active_idx=3)

    deck.section_label(s, cm(1.5), cm(6.2), "생성되는 조치")
    deck.steps(s, cm(1.5), cm(6.85), cm(14.2), [
        ("화면구성 × 1",
         "확정된 6개 패널을 벽면에 표출 — 현재 구현체: Streamlit 화면"),
        ("카메라지향 × N",
         "배치에 포함된 CCTV·열상·이동형 카메라를 해당 지점으로 지향"),
        ("지원 주체 없는 조치",
         "조용히 사라지지 않고 ‘미연동’ 결과로 남긴다"),
    ])

    deck.section_label(s, cm(17.3), cm(6.2), "실행 결과 (실측)")
    deck.bullets(s, cm(17.3), cm(6.9), cm(15.1), [
        "“북서방 무인기 2대 식별” → 조치 2건 생성",
        "미연동을 성공으로 보고하지 않는다 — 시연에서 카메라가 실제로",
        "   움직인다고 오해하게 만들지 않기 위한 설계",
    ])

    mock = deck._rect(s, cm(17.3), cm(10.0), cm(15.1), cm(6.7), MOCK_BG, radius=0.05)
    mtf = mock.text_frame
    mtf.word_wrap = True
    mtf.margin_left = mtf.margin_top = cm(0.45)
    mtf.margin_right = cm(0.3)
    rows = [
        ("조치 실행 결과 — 수행 1 · 미연동 1", MOCK_INK, True, 11),
        ("", MOCK_DIM, False, 6),
        ("● 화면구성 · 패널 6개", OK, True, 10.5),
        ("   6개 화면을 벽면에 표출했습니다.", MOCK_DIM, False, 9.5),
        ("", MOCK_DIM, False, 4),
        ("● 카메라지향 · 북서 열상감시장비(TOD)", WARN, True, 10.5),
        ("   지향 대상으로 선정됐으나 카메라 제어 장비가", MOCK_DIM, False, 9.5),
        ("   연동되지 않아 수행하지 않았습니다.", MOCK_DIM, False, 9.5),
    ]
    for i, (text, color, bold, size) in enumerate(rows):
        p = mtf.paragraphs[0] if i == 0 else mtf.add_paragraph()
        p.line_spacing = 1.3
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color

    # =================================================================
    # 슬라이드 3 — 실장비 연결
    # =================================================================
    s = deck.slide(
        "04 핵심 기술 및 구현  ·  B. 개발 과정 및 구현 결과  ·  조치 계층",
        "실장비 연결 — 구현체 하나를 추가하면 끝",
        "부대에 Video Wall·PTZ가 있으면 Actuator를 구현해 등록한다. 판단 코드는 한 줄도 바뀌지 않는다",
    )
    deck.pills(s, cm(4.75), [
        ("1", "발언 입력"), ("2", "상황 판단"), ("3", "조치 생성"),
        ("4", "장비 실행"), ("5", "결과 보고"),
    ], active_idx=4)

    deck.section_label(s, cm(1.5), cm(6.2), "확장 절차")
    deck.steps(s, cm(1.5), cm(6.85), cm(14.2), [
        ("Actuator 구현",
         "supports(kind)와 execute(action) 두 메서드만 채우면 된다"),
        ("버스 앞쪽에 등록",
         "먼저 등록된 주체가 우선 — 실장비가 스텁을 자동으로 밀어낸다"),
        ("판단 코드 수정 없음",
         "context_memory·playbook·prompts 어느 것도 건드리지 않는다"),
    ])

    deck.section_label(s, cm(17.3), cm(6.2), "검증 — 6항목 통과")
    deck.bullets(s, cm(17.3), cm(6.9), cm(15.1), [
        "가짜 PTZ 구현체 등록 → 스텁을 실제로 대체하는지 확인",
        "장비가 예외를 던져도 회의 진행이 멈추지 않는지 확인",
        "지원 주체 없는 조치가 결과로 남는지 확인",
        "미연동(skipped)과 실패(ok=False)가 구분되는지 확인",
    ])

    code = deck._rect(s, cm(17.3), cm(11.0), cm(15.1), cm(5.7), MOCK_BG, radius=0.05)
    ctf = code.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_top = cm(0.45)
    ctf.margin_right = cm(0.3)
    lines = [
        ("class 부대PTZ컨트롤러:", MOCK_INK),
        ("    name = \"PTZ 제어\"", MOCK_INK),
        ("    def supports(self, kind):", MOCK_INK),
        ("        return kind == CAMERA_AIM", MOCK_INK),
        ("    def execute(self, action):", MOCK_INK),
        ("        rest_api.aim(action.payload[\"source_id\"])", MOCK_INK),
        ("        return ActionResult(action, ok=True, ...)", MOCK_INK),
        ("", MOCK_INK),
        ("ActionBus([WallActuator(), 부대PTZ컨트롤러()])", MOCK_DIM),
    ]
    for i, (text, color) in enumerate(lines):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.line_spacing = 1.35
        r = p.add_run()
        r.text = text
        r.font.size = Pt(10)
        r.font.name = MONO
        r.font.color.rgb = color


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path,
                    default=Path(__file__).resolve().parent / "voicecue_조치계층.pptx")
    ap.add_argument("--start-page", type=int, default=24,
                    help="푸터에 찍을 쪽번호 시작값. 팀 덱에 끼워 넣을 위치에 맞춘다.")
    args = ap.parse_args()

    deck = Deck(args.start_page)
    build(deck)
    deck.prs.save(str(args.out))
    print(f"{len(deck.prs.slides.__iter__.__self__._sldIdLst)}장 -> {args.out}")


if __name__ == "__main__":
    main()
